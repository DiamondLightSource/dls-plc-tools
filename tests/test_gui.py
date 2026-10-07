"""Headless tests for the GUI's connection-mode routing and field state.

The worker thread's slots are disconnected and replaced with spies so no real
kubectl/aphyt call happens; we only assert what MainWindow emits.
"""

from types import SimpleNamespace

import pytest

# Skip (rather than error) where the Qt platform libraries are unavailable,
# e.g. a minimal container without libGL. CI runners have them.
try:
    from dls_plc_tools.gui import MainWindow
except ImportError as exc:
    pytest.skip(f"Qt unavailable: {exc}", allow_module_level=True)


@pytest.fixture
def gui(qapp):
    """A MainWindow whose request signals are diverted from the worker thread
    into spy lists, so emitting them records the call instead of running kubectl
    / connecting to a PLC."""
    win = MainWindow()
    win.request_connect.disconnect()
    win.request_connect_forward.disconnect()
    win.request_nodes.disconnect()  # don't let ticking the box run real kubectl
    direct: list = []
    forward: list = []  # tuples of (plc_ip, node, service_ip, port, cluster)
    nodes: list = []  # cluster names the Node-refresh was requested for
    win.request_connect.connect(lambda ip: direct.append(ip))
    win.request_connect_forward.connect(lambda *a: forward.append(a))
    win.request_nodes.connect(lambda c: nodes.append(c))
    yield SimpleNamespace(win=win, direct=direct, forward=forward, nodes=nodes)
    win.close()


def test_forward_fields_start_disabled(gui):
    assert not gui.win.node_edit.isEnabled()
    assert not gui.win.service_ip_edit.isEnabled()
    assert not gui.win.port_spin.isEnabled()


def test_checkbox_enables_fields(gui):
    gui.win.pfwd_check.setChecked(True)
    assert gui.win.node_edit.isEnabled()
    assert gui.win.service_ip_edit.isEnabled()
    assert gui.win.port_spin.isEnabled()
    gui.win.pfwd_check.setChecked(False)
    assert not gui.win.node_edit.isEnabled()


def test_direct_connect_when_unchecked(gui):
    gui.win.ip_edit.setText("192.168.8.10")
    gui.win._connect_clicked()
    assert gui.direct == ["192.168.8.10"]
    assert gui.forward == []


def test_forward_connect_when_checked(gui):
    gui.win.ip_edit.setText("192.168.8.10")
    gui.win.pfwd_check.setChecked(True)
    gui.win.node_edit.setCurrentText("sr08c-node.example")
    gui.win.service_ip_edit.setCurrentText("172.23.206.158")
    gui.win._connect_clicked()
    assert gui.direct == []
    assert gui.forward == [
        ("192.168.8.10", "sr08c-node.example", "172.23.206.158", 44818, "acastus")
    ]


def test_port_override_passed_through(gui):
    gui.win.ip_edit.setText("192.168.8.10")
    gui.win.pfwd_check.setChecked(True)
    gui.win.node_edit.setCurrentText("node")
    gui.win.service_ip_edit.setCurrentText("172.23.206.158")
    gui.win.port_spin.setValue(2222)
    gui.win._connect_clicked()
    assert gui.forward[0][3] == 2222


def test_forward_requires_node_and_service_ip(gui):
    gui.win.ip_edit.setText("192.168.8.10")
    gui.win.pfwd_check.setChecked(True)
    gui.win.service_ip_edit.setCurrentText("172.23.206.158")  # node left blank
    gui.win._connect_clicked()
    assert gui.forward == []
    assert gui.direct == []


def test_pod_status_led_updates(gui):
    # LED starts inactive with no pod name; each state recolours the dot.
    assert gui.win.pod_status_label.text() == "inactive"
    assert gui.win.pod_name_label.text() == ""
    seen = set()
    for state in ("pending", "good", "inactive"):
        gui.win._set_pod_status(state)
        assert gui.win.pod_status_label.text() == state
        style = gui.win.pod_led.styleSheet()
        assert "background-color" in style
        seen.add(style)
    assert len(seen) == 3  # a distinct colour per state


def test_pod_name_shown_while_active_and_cleared(gui):
    gui.win._set_pod_status("good", "pfwd-tester-44818")
    assert gui.win.pod_name_label.text() == "pfwd-tester-44818"
    gui.win._set_pod_status("inactive", "")  # teardown clears the name
    assert gui.win.pod_name_label.text() == ""


def test_ticking_box_requests_node_refresh(gui):
    # Ticking Port forward asks the worker for the cluster's live node list.
    gui.win.pfwd_check.setChecked(True)
    assert gui.nodes == ["acastus"]


def test_nodes_ready_repopulates_combo_keeping_selection(gui):
    combo = gui.win.node_edit
    typed = "sr19c-k8s-serv-01.pri.diamond.ac.uk"  # a typed/prefilled node
    combo.setCurrentText(typed)
    gui.win._on_nodes_ready(["nodeA", "nodeB"])
    items = [combo.itemText(i) for i in range(combo.count())]
    assert items == ["nodeA", "nodeB"]  # static fallback replaced by live list
    # the typed value is preserved even though it isn't in the fetched list
    assert combo.currentText() == typed


def _watch_nested(win):
    """Feed the window a flattened nested structure as if the worker had read
    it, bypassing the PLC."""
    from dls_plc_tools.structure import flatten
    from nested import make_struct

    win._current_struct = "Gauge"
    members, start = flatten(make_struct(), "Gauge")
    win._on_struct("Gauge", members, start)
    return win


def _headers(win):
    return [
        win.table.horizontalHeaderItem(c).text() for c in range(win.table.columnCount())
    ]


def _column(win, header):
    return _headers(win).index(header)


def test_nested_member_starts_collapsed_with_a_summary(gui):
    win = _watch_nested(gui.win)
    assert _headers(win) == ["Desc", "State", "▸ Ilk", "▸ Status"]
    # The cell holds a readable summary, not aphyt's CIPStructure repr.
    assert (
        win.table.item(0, _column(win, "▸ Ilk"))
        .text()
        .startswith("Trip=1 Msg=HiVac IMG={")
    )


def test_header_click_expands_and_collapses_a_nested_member(gui):
    win = _watch_nested(gui.win)
    win._header_clicked(_column(win, "▸ Ilk"))
    assert _headers(win) == [
        "Desc",
        "State",
        "▾ Ilk",
        "Ilk.Trip",
        "Ilk.Msg",
        "▸ Ilk.IMG",
        "▸ Status",
    ]
    assert win.table.item(0, _column(win, "Ilk.Msg")).text() == "HiVac"
    win._header_clicked(_column(win, "▾ Ilk"))
    assert _headers(win) == ["Desc", "State", "▸ Ilk", "▸ Status"]


def test_clicking_a_leaf_header_does_nothing(gui):
    win = _watch_nested(gui.win)
    win._header_clicked(_column(win, "Desc"))
    assert _headers(win) == ["Desc", "State", "▸ Ilk", "▸ Status"]


def test_expand_nested_checkbox_expands_every_branch(gui):
    win = _watch_nested(gui.win)
    win.expand_all.setChecked(True)
    assert _headers(win) == [
        "Desc",
        "State",
        "▾ Ilk",
        "Ilk.Trip",
        "Ilk.Msg",
        "▾ Ilk.IMG",
        "Ilk.IMG.W",
        "Ilk.IMG.B",
        "▾ Status",
        "Status.W",
        "Status.Code",
    ]
    win.expand_all.setChecked(False)
    assert _headers(win) == ["Desc", "State", "▸ Ilk", "▸ Status"]


def test_filter_matches_members_inside_a_collapsed_branch(gui):
    win = _watch_nested(gui.win)
    win.filter_edit.setText("Trip")
    # Ilk stays collapsed but is still surfaced, so the filter does not appear
    # to lose a member that exists.
    assert _headers(win) == ["▸ Ilk"]


def test_expansion_state_survives_a_refresh(gui):
    win = _watch_nested(gui.win)
    win._header_clicked(_column(win, "▸ Ilk"))
    _watch_nested(win)  # a later poll delivers a fresh read of the same tag
    assert "Ilk.Trip" in _headers(win)


def test_watching_a_different_tag_resets_expansion(gui):
    win = _watch_nested(gui.win)
    win._header_clicked(_column(win, "▸ Ilk"))
    win.struct_edit.setText("Other")
    win._watch_clicked()
    assert win._expanded == set()
    assert win._members is None


def test_in_use_filter_uses_the_marker_member(gui):
    win = _watch_nested(gui.win)
    # Desc is non-empty for two of the three declared instances.
    assert win.used_only.isChecked()
    assert win.table.rowCount() == 2
    win.used_only.setChecked(False)
    assert win.table.rowCount() == 3


def test_in_use_filter_always_keeps_index_zero(gui):
    """ether_ip assumes base-0 arrays, so index 0 is shown whenever it exists,
    even though Desc[0] is padding and blank."""
    from dls_plc_tools.structure import Member

    members = [
        Member(path="Desc", type_name="STRING", values=["", "pump1", "", "valve3"]),
        Member(path="State", type_name="INT", values=[0, 4, 0, 7]),
    ]
    win = gui.win
    win._current_struct = "Temp"
    win._on_struct("Temp", members, 0)
    assert win.used_only.isChecked()
    labels = [
        win.table.verticalHeaderItem(r).text() for r in range(win.table.rowCount())
    ]
    assert labels == ["0", "1", "3"]


def test_export_expands_nested_members_into_rows(gui, tmp_path):
    """The workbook has no expand affordance, so every nested leaf gets a row —
    and no cell holds aphyt's CIPStructure repr."""
    from types import SimpleNamespace

    from openpyxl import load_workbook

    from dls_plc_tools.gui import PlcWorker
    from nested import make_struct

    struct = make_struct()
    worker = PlcWorker()
    # A stand-in for the aphyt session: only the three calls export uses.
    worker._eip = SimpleNamespace(  # pyright: ignore[reportAttributeAccessIssue]
        update_variable_dictionary=lambda: None,
        connected_cip_dispatcher=SimpleNamespace(user_variables={"Gauge": struct}),
        read_variable=lambda name: struct,
    )
    path = tmp_path / "structs.xlsx"
    worker.export_structs(str(path))

    sheet = load_workbook(path)["Gauge"]
    rows = {r[0]: list(r[1:]) for r in sheet.iter_rows(values_only=True)}
    assert rows["Gauge"] == [1, 2, 3]  # instance indices, ARRAY[1..3]
    assert rows["Ilk.Trip"] == [1, 0, 0]  # numbers stay numeric
    # The unused third instance has an empty Msg, which reads back as a blank cell.
    assert rows["Ilk.Msg"] == ["HiVac", "Ok", None]
    assert rows["Ilk.IMG.W"][0] == "16#0041"
    # The nested struct-of-arrays spreads down the instance columns.
    assert rows["Status.W"] == ["16#0000", "16#002C", "16#0000"]
    assert not any("CIPStructure" in str(v) for row in rows.values() for v in row)


def test_export_skips_a_tag_aphyt_cannot_decode(gui, tmp_path):
    """One undecodable tag used to lose the whole workbook: aphyt raises a bare
    struct.error from its own member walk, which aborted the export. The other
    sheets are still worth having, so it is skipped and named instead."""
    import struct as _struct
    from types import SimpleNamespace

    from openpyxl import load_workbook

    from dls_plc_tools.gui import PlcWorker
    from nested import make_struct

    struct = make_struct()

    def read_variable(name):
        if name == "Temp":
            raise _struct.error("unpack requires a buffer of 2 bytes")
        return struct

    worker = PlcWorker()
    worker._eip = SimpleNamespace(  # pyright: ignore[reportAttributeAccessIssue]
        update_variable_dictionary=lambda: None,
        connected_cip_dispatcher=SimpleNamespace(
            user_variables={"Gauge": struct, "Temp": struct}
        ),
        read_variable=read_variable,
    )
    reported = []
    worker.exported.connect(reported.append)
    path = tmp_path / "structs.xlsx"
    worker.export_structs(str(path))

    assert load_workbook(path).sheetnames == ["Gauge"]
    assert "exported 1 structures" in reported[0]
    assert "could not decode: Temp (error)" in reported[0]
